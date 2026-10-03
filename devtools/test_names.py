"""Name tabs: per-tab custom sheets, settings migration, "English (Thai)" in two-language mode."""
import json, os, sys, tempfile, time
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

tmp = tempfile.mkdtemp()
os.environ["THAIW3SETUP_DATA"] = tmp  # keep the settings test away from the real profile

from core import custom
from core.custom import (COMMUNITY_ID, NAME_DOUBLE, NAME_TABS, NAME_THAI, TAB_CHARACTERS, TAB_QUESTS, TAB_SKILLS,
                         UNTRANSLATED_TAB, CustomData, Overrides, _cache_path, default_sheets, is_name_tab,
                         merged_overrides, name_label, parse_custom_xlsx, progress_of, save_custom, sheet_key)
from core.options import load_options, settings_path
from core.text_builder import _thai_for, combine
from export_names import build_tabs, classify, looks_like_name, skill_keys, write_xlsx

YEN_TH = "\u0e40\u0e22\u0e19\u0e40\u0e19\u0e40\u0e1f\u0e2d\u0e23\u0e4c"

# the workbook the exporter writes: one tab per kind, several ids per row
path = os.path.join(tmp, "names.xlsx")
write_xlsx(path, {TAB_CHARACTERS: [[[1, 2], "Yennefer", YEN_TH, ""], [[3], "Ciri", "", ""]],
                  NAME_TABS[1]: [], TAB_QUESTS: [[[9], "The Last Wish", "", "note"]]})
import openpyxl
wb = openpyxl.load_workbook(path)
ws = wb[TAB_CHARACTERS]
assert ws["D3"].value == '=IF(C3="","",B3&" ("&C3&")")', ws["D3"].value
ws["D3"] = f"Yennefer ({YEN_TH})"  # what Google exports: the computed value
wb.save(path)
data = open(path, "rb").read()
chars = parse_custom_xlsx(data, TAB_CHARACTERS)
assert chars.strings == {1: f"Yennefer ({YEN_TH})", 2: f"Yennefer ({YEN_TH})"}, chars.strings
assert chars.thai == {1: YEN_TH, 2: YEN_TH} and chars.ids == [1, 2, 3], chars
quests = parse_custom_xlsx(data, TAB_QUESTS)
assert quests.strings == {} and quests.ids == [9], quests
assert progress_of(chars) == 2 / 3 and progress_of(quests) == 0, (progress_of(chars), progress_of(quests))
assert progress_of(parse_custom_xlsx(data, NAME_TABS[1])) is None

# sheets without a THAI column list their ids too, so untranslated rows count against the progress
plain = openpyxl.Workbook()
plain.active.append(["title"])
plain.active.append(["ID", "TRANSLATE"])
plain.active.append([5, "TH"])
plain.active.append([6, None])
plain_path = os.path.join(tmp, "plain.xlsx")
plain.save(plain_path)
plain_data = parse_custom_xlsx(open(plain_path, "rb").read())
assert plain_data.ids == [5, 6] and plain_data.strings == {5: "TH"} and progress_of(plain_data) == 0.5, plain_data

# a tab that marks rows green: green rows are done (even left blank on purpose), yellow / white / unfilled
# rows are not, rows of any other colour are left out
from openpyxl.styles import PatternFill
green = openpyxl.Workbook()
gws = green.active
gws.append(["title"])
gws.append(["ID", "ENGLISH", "TRANSLATE"])
colors = [("FF00FF00", "TH"), ("FFB7E1CD", None), ("FFFFFF00", "TH?"), ("FFFFF2CC", None), ("FFFFFFFF", None),
          (None, None), ("FF999999", None), ("FFFF9900", "TH?"), ("FF9FC5E8", "TH"), ("FFFF0000", None)]
for sid, (color, text) in enumerate(colors, 1):
    gws.append([sid, f"en{sid}", text])
    if color:
        gws[f"C{sid + 2}"].fill = PatternFill("solid", fgColor=color)
green_path = os.path.join(tmp, "green.xlsx")
green.save(green_path)
green_data = parse_custom_xlsx(open(green_path, "rb").read())
assert green_data.done == [1, 2] and green_data.skipped == [7, 8, 9, 10], green_data
assert progress_of(green_data) == 2 / 6, progress_of(green_data)
from pathlib import Path
green_cache = Path(tmp) / "green.json.gz"
save_custom(green_cache, green_data, 0)
cached = custom.load_custom(green_cache)[0]
assert cached.done == [1, 2] and cached.skipped == [7, 8, 9, 10], cached
try:
    parse_custom_xlsx(data, "missing")
    raise AssertionError("missing tab accepted")
except ValueError:
    pass

# tabs of one sheet get their own keys and caches; the old untranslated cache file stays in use
keys = [sheet_key(s) for s in default_sheets()]
assert len(keys) == len(set(keys)), keys
assert _cache_path(COMMUNITY_ID, UNTRANSLATED_TAB) == _cache_path(COMMUNITY_ID)
assert len({_cache_path(COMMUNITY_ID, t) for t in NAME_TABS}) == len(NAME_TABS)

# name modes: Thai only, "English (Thai)", or English for a switched-off tab even if the main sheet has Thai
now = time.time()
for tab in NAME_TABS:
    save_custom(_cache_path(COMMUNITY_ID, tab), {TAB_CHARACTERS: chars, TAB_QUESTS: quests}.get(tab, CustomData("", {})), now)
name_sheets = [s for s in default_sheets() if is_name_tab(s)]


def overrides_for(mode: str, enabled: bool) -> Overrides:
    sheets = [dict(s, enabled=enabled, name_mode=mode) for s in name_sheets]
    return merged_overrides(sheets, progress=lambda f, m: None)


ov = overrides_for(NAME_THAI, True)
assert ov.strings == {1: YEN_TH, 2: YEN_TH} and ov.keep_english == {3, 9} and ov.plain == {1, 2, 3, 9}, ov
ov = overrides_for(NAME_DOUBLE, True)
assert ov.strings == {1: f"Yennefer ({YEN_TH})", 2: f"Yennefer ({YEN_TH})"} and ov.keep_english == {3, 9}, ov
ov = overrides_for(NAME_DOUBLE, False)
assert ov.strings == {} and ov.keep_english == {1, 2, 3, 9}, ov
assert _thai_for(1, "Yennefer", {1: YEN_TH}, None, ov) == "Yennefer"

# settings from 0.3.0 (bare community id, no tab) gain the name tabs once, disabled
old = [dict(s) for s in default_sheets() if not s.get("tab") or s["tab"] == UNTRANSLATED_TAB]
for s in old:
    s.pop("tab", None)
    s.pop("gid", None)
settings_path().write_text(json.dumps({"custom_sheets": old,
                                       "known_default_sheets": [s["sheet_id"] for s in old]}), encoding="utf-8")
opts = load_options()
community = [s for s in opts.custom_sheets if s["sheet_id"] == COMMUNITY_ID]
assert [s["tab"] for s in community] == [UNTRANSLATED_TAB, *NAME_TABS], community
assert community[0]["enabled"] and not any(s["enabled"] for s in community[1:])
opts.custom_sheets = [s for s in opts.custom_sheets if s.get("tab") != TAB_QUESTS]
settings_path().write_text(json.dumps({"custom_sheets": opts.custom_sheets,
                                       "known_default_sheets": opts.known_default_sheets}), encoding="utf-8")
assert TAB_QUESTS not in [s.get("tab") for s in load_options().custom_sheets]  # removal sticks

# settings from 0.3.1 named the name tabs after the worksheet; they gain the label prefix
for s in opts.custom_sheets:
    if s.get("tab") in NAME_TABS:
        s["name"] = s["tab"]
settings_path().write_text(json.dumps({"custom_sheets": opts.custom_sheets,
                                       "known_default_sheets": opts.known_default_sheets}), encoding="utf-8")
loaded = [s for s in load_options().custom_sheets if s.get("tab") in NAME_TABS]
assert {s["tab"]: s["name"] for s in loaded} == {t: name_label(t) for t in NAME_TABS if t != TAB_QUESTS}, loaded

# settings without a name mode default to "English (Thai)"
for s in opts.custom_sheets:
    s.pop("name_mode", None)
settings_path().write_text(json.dumps({"custom_sheets": opts.custom_sheets,
                                       "known_default_sheets": opts.known_default_sheets}), encoding="utf-8")
assert all(s["name_mode"] == NAME_DOUBLE for s in load_options().custom_sheets if is_name_tab(s))

# two-language mode does not repeat the English name
assert combine(f"Yennefer ({YEN_TH})", "Yennefer", True) == f"Yennefer ({YEN_TH})"
assert combine(YEN_TH, "Yennefer", True) == f"{YEN_TH}  [Yennefer]"

# re-running the exporter keeps typed Thai, notes and hand-moved rows
found = {TAB_CHARACTERS: {"Yennefer": [1], "Novigrad": [5]}, NAME_TABS[1]: {}, TAB_QUESTS: {}}
existing = {NAME_TABS[1]: {"Novigrad": [[5], "TH", ""]}, TAB_CHARACTERS: {"Gone": [[7], "", "keep"]}}
tabs = build_tabs(found, {}, existing)
assert tabs[NAME_TABS[1]] == [[[5], "Novigrad", "TH", ""]], tabs
assert tabs[TAB_CHARACTERS] == [[[7], "Gone", "", "keep"], [[1], "Yennefer", "", ""]], tabs

# skills move to their tab even from a tab someone put them in, keeping the Thai
tabs = build_tabs({TAB_SKILLS: {"Whirl": [4, 8]}}, {}, {TAB_CHARACTERS: {"Whirl": [[4], "TH", ""]}})
assert tabs[TAB_SKILLS] == [[[4, 8], "Whirl", "TH", ""]] and tabs[TAB_CHARACTERS] == [], tabs
assert "skill_name_sword_s1" in skill_keys() and "skill_name_mutation_1" in skill_keys()

# each id goes to the tab of what the game uses it for
from core.custom import TAB_GWENT, TAB_ITEMS, TAB_MONSTERS, TAB_OTHER, TAB_PLACES
from core.w3strings import hash_key
from export_names import category, evidence
from game_index import GameIndex

index = GameIndex(
    key_names={hash_key("gwint_name_triss"): ("gwint_name_triss", ["def_gwint_cards_final.xml"]),
               hash_key("item_name_harpy"): ("item_name_harpy", ["def_item_crafting_weapons.xml"]),
               hash_key("map_location_lake_village_ft"): ("map_location_lake_village_ft", [".w2em"]),
               hash_key("aard"): ("aard", ["geralt_skills.xml", ".w2scene"]),
               hash_key("ea_explorer"): ("ea_explorer", ["achievements.xml"]),
               hash_key("blanka"): ("blanka", [".w2scene"])},
    journal={11: ["characters"], 12: ["bestiary", "quests"], 13: ["quests"], 14: ["tutorial"]},
    entities={21: ["quests\\main_npcs\\yennefer.w2ent"], 22: ["characters\\npc_entities\\monsters\\hag_grave.w2ent"],
              23: ["quests\\part_2\\quest_files\\q107_swamps\\characters\\q107_monster_anna.w2ent"],
              24: ["quests\\minor_quests\\mq1051_monster_hunt_nilfgaard1\\characters\\mq1051_nilfgaard_officer.w2ent"],
              25: ["characters\\npc_entities\\secondary_npc\\anna.w2ent",
                   "characters\\npc_entities\\monsters\\hag_grave_lvl1__barons_wife.w2ent"]})
assert category(1, hash_key("gwint_name_triss"), index) == TAB_GWENT
assert category(2, hash_key("item_name_harpy"), index) == TAB_ITEMS
assert category(3, hash_key("map_location_lake_village_ft"), index) == TAB_PLACES
assert category(4, hash_key("aard"), index) == TAB_SKILLS
assert category(5, hash_key("ea_explorer"), index) == TAB_OTHER
assert category(6, hash_key("blanka"), index) == TAB_CHARACTERS
assert category(11, None, index) == TAB_CHARACTERS and category(12, None, index) == TAB_MONSTERS
assert category(13, None, index) is None and category(14, None, index) == TAB_OTHER
assert category(21, None, index) == TAB_CHARACTERS
assert category(22, None, index) == TAB_MONSTERS and category(23, None, index) == TAB_MONSTERS
# a contract giver in a monster hunt quest, and a person who also has a monster form
assert category(24, None, index) == TAB_CHARACTERS and evidence(25, None, index) == (TAB_MONSTERS, False)
assert category(99, 12345, index) is None
# the Gwent key wins over the journal
assert category(11, hash_key("gwint_name_triss"), index) == TAB_GWENT

# a row whose ids the game uses for different things is split, every part keeping the Thai and the note
TRISS_TH = "\u0e17\u0e23\u0e34\u0e2a"
existing = {TAB_CHARACTERS: {"Triss Merigold": [[1061858, 1065275], TRISS_TH, "note"],
                             "Nobody": [[30], "", ""], "Moved": [[31], "", ""]},
            TAB_QUESTS: {"Contract: Imp": [[40], "", ""]}, NAME_TABS[1]: {"Moved": [[32], "", ""]}}
found = {TAB_CHARACTERS: {"Triss Merigold": [1061858, 1065275], "Nobody": [30], "Moved": [31]},
         TAB_QUESTS: {"Contract: Imp": [40]}}
known = {1061858: TAB_CHARACTERS, 1065275: TAB_GWENT, 40: TAB_ITEMS}
tabs = build_tabs(found, {}, existing, known)
assert [[1061858], "Triss Merigold", TRISS_TH, "note"] in tabs[TAB_CHARACTERS], tabs
# a NOTE naming the tab of the other half of a split row is dropped from this half only
GWENT_TH = "\u0e01\u0e32\u0e23\u0e4c\u0e14\u0e40\u0e01\u0e27\u0e19\u0e15\u0e4c"
split = build_tabs({}, {}, {TAB_CHARACTERS: {"Triss Merigold": [[1061858, 1065275], TRISS_TH, GWENT_TH],
                                             "Bart": [[1059068], "", "\u0e21\u0e2d\u0e19\u0e2a\u0e40\u0e15\u0e2d\u0e23\u0e4c"]}},
                   {**known, 1059068: TAB_CHARACTERS})
assert [[1061858], "Triss Merigold", TRISS_TH, ""] in split[TAB_CHARACTERS], split
assert split[TAB_GWENT] == [[[1065275], "Triss Merigold", TRISS_TH, GWENT_TH]], split
assert [[1059068], "Bart", "", "\u0e21\u0e2d\u0e19\u0e2a\u0e40\u0e15\u0e2d\u0e23\u0e4c"] in split[TAB_CHARACTERS], split
assert tabs[TAB_GWENT] == [[[1065275], "Triss Merigold", TRISS_TH, "note"]], tabs
# ids the game does not explain stay where they are; quests keep their tab
assert [[30], "Nobody", "", ""] in tabs[TAB_CHARACTERS] and [[31], "Moved", "", ""] in tabs[TAB_CHARACTERS], tabs
assert tabs[TAB_QUESTS] == [[[40], "Contract: Imp", "", ""]] and tabs[TAB_ITEMS] == [], tabs
# unexplained ids join the other ids of their name when those agree on one tab, otherwise stay
tabs = build_tabs({TAB_CHARACTERS: {"Aether": [50, 51], "Dandelion": [60, 61, 62]}}, {}, {},
                  {51: TAB_ITEMS, 61: TAB_CHARACTERS, 62: TAB_GWENT})
assert tabs[TAB_ITEMS] == [[[50, 51], "Aether", "", ""]], tabs
assert tabs[TAB_CHARACTERS] == [[[60, 61], "Dandelion", "", ""]] and tabs[TAB_GWENT] == [[[62], "Dandelion", "", ""]], tabs
# a weak character answer (a plain entity) follows strong siblings, and keeps its tab when alone
tabs = build_tabs({TAB_CHARACTERS: {"Ghoul": [70, 71], "Molly": [72]}}, {}, {},
                  {70: TAB_CHARACTERS, 71: TAB_MONSTERS, 72: TAB_CHARACTERS}, {70, 72})
assert tabs[TAB_MONSTERS] == [[[70, 71], "Ghoul", "", ""]] and tabs[TAB_CHARACTERS] == [[[72], "Molly", "", ""]], tabs
# with no explained id in the rows, the same text elsewhere in the game decides
tabs = build_tabs({TAB_CHARACTERS: {"Aerondight": [80], "Fiend": [81, 82]}}, {}, {}, {82: TAB_GWENT}, set(),
                  {"Aerondight": {TAB_ITEMS}, "Fiend": {TAB_CHARACTERS}})
assert tabs[TAB_ITEMS] == [[[80], "Aerondight", "", ""]] and tabs[TAB_GWENT] == [[[81, 82], "Fiend", "", ""]], tabs
# ids the game says nothing about go where their NOTE says
tabs = build_tabs({TAB_CHARACTERS: {"Dantan Glade": [90]}}, {},
                  {TAB_CHARACTERS: {"Dantan Glade": [[90], "", "\u0e2a\u0e16\u0e32\u0e19\u0e17\u0e35\u0e48"]}}, {})
assert tabs[TAB_PLACES] == [[[90], "Dantan Glade", "", "\u0e2a\u0e16\u0e32\u0e19\u0e17\u0e35\u0e48"]], tabs
# character rows the journal knows get a NOTE, over a tab word but not over free text
from export_names import JOURNAL_NOTE, mark_journal
MONSTER_TH = "\u0e21\u0e2d\u0e19\u0e2a\u0e40\u0e15\u0e2d\u0e23\u0e4c"
marked = {TAB_CHARACTERS: [[[1], "A", "", ""], [[2], "B", "", MONSTER_TH], [[3], "C", "", "ask Bob"], [[4], "D", "", ""]]}
mark_journal(marked, {1, 2, 3})
assert [r[3] for r in marked[TAB_CHARACTERS]] == [JOURNAL_NOTE, JOURNAL_NOTE, "ask Bob", ""], marked
# hand-checked evidence replaces the NOTE of the row holding the id
from export_names import HAND_NOTES, hand_notes
noted = {TAB_CHARACTERS: [[[5, 6], "E", "", MONSTER_TH], [[7], "F", "", "x"]]}
hand_notes(noted, {6: "evidence"})
assert [r[3] for r in noted[TAB_CHARACTERS]] == ["evidence", "x"], noted
import json
assert all(k.isdigit() for k in json.loads(HAND_NOTES.read_text(encoding="utf-8")))
# speaker names count before entities, and are strong
voiced = GameIndex({hash_key("anna"): ("anna", ["scene_voice_tags.csv"])}, {}, {26: ["x\\anna.w2ent"]})
assert evidence(26, hash_key("anna"), voiced) == (TAB_CHARACTERS, True)
assert evidence(21, None, index) == (TAB_CHARACTERS, False) and evidence(22, None, index) == (TAB_MONSTERS, True)
# but monster entities count before speaker names, which voice monsters too
growl = GameIndex({hash_key("cyclops"): ("cyclops", ["scene_voice_tags.csv"])}, {},
                  {27: ["characters\\npc_entities\\monsters\\cyclop_lvl1.w2ent"]})
assert evidence(27, hash_key("cyclops"), growl) == (TAB_MONSTERS, True)
owl = GameIndex({hash_key("molly"): ("molly", ["scene_voice_tags.csv"])}, {},
                {28: ["characters\\npc_entities\\animals\\owl.w2ent"]})
assert evidence(28, hash_key("molly"), owl) == (TAB_CHARACTERS, True)
# an NPC (weak) does not follow the Gwent card named after it
tabs = build_tabs({TAB_CHARACTERS: {"Vlodimir": [91, 92]}}, {}, {}, {91: TAB_CHARACTERS, 92: TAB_GWENT}, {91})
assert tabs[TAB_CHARACTERS] == [[[91], "Vlodimir", "", ""]] and tabs[TAB_GWENT] == [[[92], "Vlodimir", "", ""]], tabs
# nor the tutorial page named after it, in the rows or elsewhere in the game
tabs = build_tabs({TAB_OTHER: {"Runewright": [93]}}, {}, {}, {93: TAB_CHARACTERS}, {93}, {"Runewright": {TAB_OTHER}})
assert tabs[TAB_CHARACTERS] == [[[93], "Runewright", "", ""]], tabs
# filled rows no longer found still follow the game files
tabs = build_tabs({}, {}, {TAB_CHARACTERS: {"Triss Merigold": [[1061858, 1065275], TRISS_TH, ""]}}, known)
assert tabs[TAB_GWENT] == [[[1065275], "Triss Merigold", TRISS_TH, ""]], tabs

assert looks_like_name("Olgierd von Everec") and not looks_like_name("LEAD QA") and not looks_like_name("VSync")
assert classify("Nilfgaardian Armor Set", set(), False, set()) is None
assert classify("Crane Isle", set(), False, set()) == NAME_TABS[1]
assert classify("Food", set(), False, {"food"}) is None
# old backgrounds are cleared down to the last old row; only the characters THAI column is painted
from export_names import _fill_requests
fills = _fill_requests(9, TAB_CHARACTERS, 3, 100)
assert fills[0]["repeatCell"]["range"]["endRowIndex"] == 100
assert fills[1]["repeatCell"]["range"] == {"sheetId": 9, "startRowIndex": 2, "endRowIndex": 5,
                                           "startColumnIndex": 2, "endColumnIndex": 3}
assert len(_fill_requests(9, TAB_PLACES, 3, 100)) == 1

print("test_names ok")
