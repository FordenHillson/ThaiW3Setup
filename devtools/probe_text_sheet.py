"""Coverage of an English->Thai text sheet against the game's en.w3strings, vs the id-based sheets."""
import os, sys
from collections import Counter, defaultdict
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from gamepath import GAME
import openpyxl
from core.game_detect import identify
from core.sheet import get_translations
from core.text_builder import _load_merged

XLSX = sys.argv[1]


def norm(s):
    return " ".join(str(s).split())


wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
pairs = defaultdict(Counter)
per_tab = {}
for name in wb.sheetnames[1:]:
    n = 0
    for row in wb[name].iter_rows(min_row=5, max_col=3, values_only=True):
        if not row or row[0] is None or len(row) < 3 or row[2] is None:
            continue
        en, th = norm(row[0]), str(row[2]).strip()
        if en and th and th != en or (en and th):
            pairs[en][th] += 1
            n += 1
    per_tab[name] = n
print("rows per tab", per_tab)
conflicts = sum(1 for v in pairs.values() if len(v) > 1)
print("unique english", len(pairs), "with conflicting thai", conflicts)

english = _load_merged(identify(GAME), "en")
ids = [(sid, s) for sid, s in english.strings.items() if s.strip()]
tr = get_translations(allow_online=False)
by_text = sum(1 for sid, s in ids if norm(s) in pairs)
by_id = sum(1 for sid, s in ids if sid in tr.thai)
union = sum(1 for sid, s in ids if sid in tr.thai or norm(s) in pairs)
untrans_in_sheet = sum(1 for sid, s in ids if norm(s) in pairs and pairs[norm(s)].most_common(1)[0][0].strip() == s.strip())
print(f"non-empty game strings {len(ids)}")
print(f"text sheet matches {by_text} ({100*by_text/len(ids):.2f}%), of which thai == english {untrans_in_sheet}")
print(f"id sheets        {by_id} ({100*by_id/len(ids):.2f}%)")
print(f"union            {union} ({100*union/len(ids):.2f}%)")
missing = [s for sid, s in ids if sid not in tr.thai and norm(s) not in pairs]
print("sample still missing:", [ascii(m[:60]) for m in missing[:8]])
